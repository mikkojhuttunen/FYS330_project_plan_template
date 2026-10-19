#!/usr/bin/env python3
"""
gantt_xlsx_to_tex.py -- regenerate the FYS.330 project-plan Gantt table in a
.tex file from the FYS330 Gantt workbook (sheets: Weeks, Tasks).

Only the *editable* input columns are read (formulas / cached values are
ignored, so a stale or never-recalculated workbook is fine):
  Weeks: Calendar Week | Week Start | Week End | Phase | Is Break (Y/N) | Break Label
  Tasks: ID | Task | Responsible | Type | Start (Cal Wk) | End (Cal Wk) | Deliverable
         optional extra column: "Gantt label"  (short name used in the chart;
         falls back to the Task column)

Usage:
  python gantt_xlsx_to_tex.py book.xlsx Chapters/ChapII_Implementation.tex          # update in place
  python gantt_xlsx_to_tex.py book.xlsx file.tex --out new.tex                       # write elsewhere
  python gantt_xlsx_to_tex.py book.xlsx file.tex --leaders "Planning=A,Reporting=B"  # phase leaders

The tabular is replaced between the markers
  % BEGIN GANTT (auto-generated ...)   /   % END GANTT
On the first run (no markers yet) they are inserted around the tabular that
follows \\label{tab:gantt} (and the active \\setlength{\\wkcolwidth}... line).
Phase leaders, e.g. "Planning (Team lead A)", are kept from the existing header
unless overridden with --leaders.
"""
import argparse, re, sys, warnings
from collections import OrderedDict
import openpyxl

warnings.filterwarnings("ignore")
BEGIN = "% BEGIN GANTT (auto-generated from the Gantt workbook -- edit the .xlsx, not this block)"
END = "% END GANTT"
LABEL = r"\label{tab:gantt}"


# ----------------------------------------------------------------- reading
def tex_escape(s):
    rep = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#",
           "_": r"\_", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}
    return "".join(rep.get(c, c) for c in str(s))


def read_book(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    warns = []
    # ---- Weeks
    ws = wb["Weeks"]
    hdr = [str(c.value).strip() if c.value else "" for c in ws[1]]
    ix = {h: i for i, h in enumerate(hdr)}
    weeks, pw = [], 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        brk = str(row[ix["Is Break (Y/N)"]]).strip().upper() == "Y"
        if not brk:
            pw += 1
        weeks.append(dict(cal=str(row[0]).strip(), phase=str(row[ix["Phase"]]).strip(), brk=brk,
                          label=(row[ix["Break Label"]] or "").strip() if brk else "",
                          pw=None if brk else pw))
    # ---- Tasks
    ws = wb["Tasks"]
    hdr = [str(c.value).strip() if c.value else "" for c in ws[1]]
    ix = {h: i for i, h in enumerate(hdr)}
    tasks, seen = [], {}
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0]:
            continue
        t = dict(id=str(row[ix["ID"]]).strip(), name=str(row[ix["Task"]]).strip(),
                 resp=str(row[ix["Responsible"]] or "").strip(), type=str(row[ix["Type"]]).strip(),
                 start=str(row[ix["Start (Cal Wk)"]]).strip(), end=str(row[ix["End (Cal Wk)"]]).strip(), row=r)
        if "Gantt label" in ix and row[ix["Gantt label"]]:
            t["name"] = str(row[ix["Gantt label"]]).strip()
        if t["id"] in seen:
            o = seen[t["id"]]
            same = all(o[k] == t[k] for k in ("name", "resp", "type", "start", "end"))
            if same:
                warns.append(f"Tasks row {r}: '{t['id']}' duplicates row {o['row']} (identical) -- skipped")
                continue
            sys.exit(f"ERROR: Tasks row {r}: ID '{t['id']}' conflicts with row {o['row']}")
        seen[t["id"]] = t
        tasks.append(t)
    return weeks, tasks, warns


# ---------------------------------------------------------------- building
def build_layout(weeks):
    """Group the calendar spine into segments: work spans (per phase) and break blocks."""
    segs = []
    for i, w in enumerate(weeks):
        key = ("break", w["label"]) if w["brk"] else ("work", w["phase"])
        if segs and segs[-1]["key"] == key:
            segs[-1]["weeks"].append(i)
        else:
            segs.append(dict(key=key, weeks=[i]))
    return segs


def generate(weeks, tasks, leaders, trail="0.25cm"):
    warns = []
    wk_index = {w["cal"]: i for i, w in enumerate(weeks)}
    segs = build_layout(weeks)
    nwork = sum(1 for w in weeks if not w["brk"])

    # slots: one per column after the two label columns
    slot_of_week, spec, hdr1, hdr_pw, hdr_cal = {}, [], [], [], []
    slot = 0
    grey_w = 0.0
    prev = None
    nslots = 0
    for si, s in enumerate(segs):
        kind, name = s["key"]
        last = si == len(segs) - 1
        if kind == "work":
            if prev == "work":
                spec.append("|")  # phase boundary without a break
            n = len(s["weeks"])
            spec.append(r"*{%d}{>{\centering\arraybackslash}p{\wkcolwidth}}" % n)
            for wi in s["weeks"]:
                slot_of_week[wi] = slot
                slot += 1
                hdr_pw.append(r"{\tiny %d}" % weeks[wi]["pw"])
                hdr_cal.append(r"{\tiny %d}" % int(weeks[wi]["cal"][1:3]))
        else:
            if last:  # trailing break: spacer + marker column (room for a deadline marker)
                spec.append(r"|>{\columncolor{gray!35}}p{0.01cm}>{\columncolor{gray!35}}p{%s}|@{}" % trail)
                hdr_pw += ["", ""]; hdr_cal += ["", ""]
                for wi in s["weeks"]:
                    slot_of_week[wi] = slot + 1
                slot += 2; grey_w += 0.01 + float(trail[:-2])
            else:
                spec.append(r"|>{\columncolor{gray!35}}p{0.1cm}|")
                hdr_pw.append(""); hdr_cal.append("")
                for wi in s["weeks"]:
                    slot_of_week[wi] = slot
                slot += 1; grey_w += 0.1
        prev = kind
    nslots = slot
    # top header row: one multicolumn per phase; a phase interrupted by an interior break is merged across it
    def nsl(sg, k):
        return len(sg["weeks"]) if sg["key"][0] == "work" else (2 if k == len(segs) - 1 else 1)
    k = 0
    while k < len(segs):
        sg = segs[k]
        if sg["key"][0] != "work":
            hdr1 += [""] * nsl(sg, k); k += 1; continue
        j = k
        while j + 2 < len(segs) and segs[j + 1]["key"][0] == "break" and j + 1 < len(segs) - 1 \
                and segs[j + 2]["key"] == sg["key"]:
            j += 2
        n = sum(nsl(segs[m], m) for m in range(k, j + 1))
        name = sg["key"][1]
        lab = leaders.get(name, name)
        lab = tex_escape(lab) if lab == name else lab
        lab = (r"\makebox[0pt]{\textbf{%s}}" if n < 6 else r"\textbf{%s}") % lab
        hdr1.append(r"\multicolumn{%d}{c}{%s}" % (n, lab))
        k = j + 1
    if segs[-1]["key"][0] == "work":
        spec.append("@{}")

    ncols = 2 + nslots
    idw = "0.35cm" if max(len(t["id"]) for t in tasks) <= 2 else "0.55cm"  # room for IDs like T10
    colspec = r"@{}p{%s} p{3.6cm}" % idw + "\n   " + "\n   ".join(spec)
    nrules = "".join(spec).count("|")  # each vertical rule eats a little width
    slack = round(0.26 / 6 * nrules, 2)

    rows = []
    for t in tasks:
        for k in ("start", "end"):
            if t[k] not in wk_index:
                sys.exit(f"ERROR: {t['id']}: {k} week '{t[k]}' not on Weeks sheet")
        a, b = wk_index[t["start"]], wk_index[t["end"]]
        if a > b:
            sys.exit(f"ERROR: {t['id']}: start after end")
        cells = [""] * nslots
        for wi in range(a, b + 1):
            w = weeks[wi]
            if t["type"] == "Deliverable":
                if wi == b:  # a deliverable is a single marker in its (end) week
                    cells[slot_of_week[wi]] = r"\gdl{%s}" % t["id"]
            elif w["brk"]:
                if wi in (a, b):
                    warns.append(f"{t['id']}: starts/ends in break week {w['cal']} -- no cell drawn")
            else:
                macro = "gwkm" if t["type"] == "Milestone" else "gwk"
                cells[slot_of_week[wi]] = r"\%s{%d}" % (macro, w["pw"])
        resp = t["resp"]
        if t["type"] == "Milestone":
            resp = "milestone, " + (resp[0].lower() + resp[1:] if resp.lower() in ("whole team", "individual", "team") else resp)
        label = "%s \\textit{(%s)}" % (tex_escape(t["name"]), tex_escape(resp)) if resp else tex_escape(t["name"])
        rows.append("%s & %s\n   & %s \\\\" % (tex_escape(t["id"]), label, " & ".join(cells)))

    tab = ["\\setlength{\\wkcolwidth}{\\dimexpr(\\textwidth-0.5cm-4cm-%.2fcm-%d\\tabcolsep)/%d\\relax}"
           % (grey_w + slack, 2 * (ncols - 1), nwork),
           "\\begin{tabular}{%s}" % colspec, "\\toprule",
           "\\textbf{ID} & \\textbf{Task (responsible)}\n   & " + " & ".join(hdr1) + " \\\\",
           "& \\textit{Project wk} & " + " & ".join(hdr_pw) + " \\\\",
           "& \\textit{Calendar wk} & " + " & ".join(hdr_cal) + " \\\\",
           "\\midrule"] + rows + ["\\bottomrule", "\\end{tabular}"]
    return "\n".join(tab), warns, ncols


# ------------------------------------------------------------- tex patching
def existing_leaders(block):
    d = OrderedDict()
    block = re.sub(r"(?<!\\)%[^\n]*", "", block)  # ignore commented-out header lines
    for lab in re.findall(r"\\multicolumn\{\d+\}\{c\}\{(?:\\makebox\[0pt\]\{)?\\textbf\{([^}]*)\}", block):
        d[lab.split(" (")[0]] = lab
    return d


def patch(tex, new_block):
    if BEGIN in tex:
        a, b = tex.index(BEGIN), tex.index(END)
        return tex[:a] + BEGIN + "\n" + new_block + "\n" + tex[b:]
    i = tex.index(LABEL)
    t0 = tex.index("\\begin{tabular}", i)
    t1 = tex.index("\\end{tabular}", t0) + len("\\end{tabular}")
    # include the active \setlength{\wkcolwidth} line directly above the tabular
    pre = tex[i:t0]
    m = list(re.finditer(r"^\\setlength\{\\wkcolwidth\}.*\n", pre, re.M))
    start = i + m[-1].start() if m else t0
    return tex[:start] + BEGIN + "\n" + new_block + "\n" + END + tex[t1:]


def current_block(tex):
    if BEGIN in tex:
        return tex[tex.index(BEGIN):tex.index(END)]
    i = tex.index(LABEL)
    t0 = tex.index("\\begin{tabular}", i)
    return tex[t0:tex.index("\\end{tabular}", t0)]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx"); ap.add_argument("tex")
    ap.add_argument("--out"); ap.add_argument("--leaders", default="")
    ap.add_argument("--trail-width", default="0.25cm", help="width of the last grey marker column (D5 needs ~0.4cm)")
    a = ap.parse_args()
    weeks, tasks, w1 = read_book(a.xlsx)
    tex = open(a.tex, encoding="utf-8").read()
    leaders = existing_leaders(current_block(tex))
    for kv in filter(None, a.leaders.split(",")):
        k, v = kv.split("=", 1)
        leaders[k.strip()] = f"{k.strip()} ({v.strip()})"
    block, w2, ncols = generate(weeks, tasks, leaders, a.trail_width)
    open(a.out or a.tex, "w", encoding="utf-8").write(patch(tex, block))
    for w in w1 + w2:
        print("WARNING:", w, file=sys.stderr)
    print(f"OK: {len(tasks)} tasks, {sum(not w['brk'] for w in weeks)} project weeks, "
          f"{len(weeks)} calendar weeks, {ncols} table columns -> {a.out or a.tex}")


if __name__ == "__main__":
    main()
